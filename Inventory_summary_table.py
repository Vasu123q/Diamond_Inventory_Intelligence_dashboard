import pandas as pd
import numpy as np
from openpyxl import load_workbook
from pathlib import Path

project_root = Path(__file__).resolve().parent

# =========================
# LOAD FILES
# =========================
df1 = pd.read_excel(project_root/'L1'/'ML_Prediction_Level1.xlsx')
df2 = pd.read_excel(project_root/'L2'/'ML_Prediction_Level2.xlsx')
df3 = pd.read_excel(project_root/'L3'/'ML_Prediction_Level3.xlsx')

dataframes = [
    ("L1", df1),
    ("L2", df2),
    ("L3", df3)
]

# =========================
# COMMON SETTINGS
# =========================
available_status = ["Consignment In","Memo In","Available to Use","In Stock"]
memo_out_status = ["Consignment Out","Memo Out"]

locations = ["HK","USA","DUBAI","IND"]

col_map = {
    "Total Stones": 3,
    "Memo Out": 4,
    "AVL Stock": 5,
    "Unpublish AVL(Not On Rap)": 6,
    "On Hand": 7,
    "On Hold": 8,
    "For Web": 9,
    "Transit": 10,
    "Reserved": 11,
    "Consume": 12,
    "Under Certification": 13
}

# =========================
# PROCESS EACH LEVEL
# =========================
for name, df in dataframes:

    print(f"Processing {name}...")

    # ===== CLEAN =====
    df['Stone Location'] = df['Stone Location'].replace({
        'UAE':'DUBAI',
        'INDIA':'IND'
    })

    df['Stone Location'] = df['Stone Location'].astype(str).str.upper().str.strip()
    df['Status'] = df['Status'].astype(str).str.strip()
    df['ItemCD'] = df['ItemCD'].astype(str).str.strip()
    df['Lab'] = df['Lab'].astype(str).str.strip()

    # ===== SUMMARY =====
    summary = {}

    for loc, g in df.groupby('Stone Location'):

        summary[loc] = {
            "Total Stones": len(g),
            "Memo Out": g[g['Status'].isin(memo_out_status)].shape[0],
            "AVL Stock": g[g['Status'].isin(available_status)].shape[0],
            "Unpublish AVL(Not On Rap)": g[
                (g['Status'].isin(available_status)) & (g['Publish']==False)
            ].shape[0],
            "On Hand": g[
                (g['Status'].isin(available_status)) &
                (g['Publish']==False) &
                (~g['ItemCD'].str.startswith(('FE','MO','KV'))) &
                (~g['Lab'].isin(['NC','',None]))
            ].shape[0],
            "On Hold": g[g['Status']=="On Hold"].shape[0],
            "For Web": g[
                (g['Status'].isin(available_status)) & (g['Publish']==True)
            ].shape[0],
            "Transit": g[g['Status']=="Transit"].shape[0],
            "Reserved": g[g['Status']=="Reserved"].shape[0],
            "Consume": g[g['Status']=="Consume"].shape[0],
            "Under Certification": g[g['Status']=="Under Certification"].shape[0],
        }

    # ===== LOAD TEMPLATE =====
    wb = load_workbook(project_root / "Inventory_SUMMARY_Table.xlsx")
    ws = wb.active

    # ===== WRITE COUNTS =====
    start_row = 4

    for i, loc in enumerate(locations):
        row = start_row + i
        data = summary.get(loc, {k:0 for k in col_map})

        for key, col in col_map.items():
            ws.cell(row=row, column=col).value = data[key]

    # ===== TOTAL ROW =====
    total_row = start_row + len(locations)

    for key, col in col_map.items():
        total_val = sum(summary.get(loc, {}).get(key, 0) for loc in locations)
        ws.cell(row=total_row, column=col).value = total_val

    # ===== RATIO TABLE =====
    ratio_start = 14

    for i, loc in enumerate(locations):
        row = ratio_start + i
        data = summary.get(loc, {k:0 for k in col_map})
        total = data.get("Total Stones", 0)

        for key, col in col_map.items():

            value = data[key]

            if key == "Total Stones":
                ws.cell(row=row, column=col).value = value
            else:
                ratio = (value / total) if total != 0 else 0
                cell = ws.cell(row=row, column=col)
                cell.value = ratio
                cell.number_format = '0%'

    # ===== TOTAL RATIO =====
    ratio_total_row = ratio_start + len(locations)
    total_stones = ws.cell(row=total_row, column=3).value

    for key, col in col_map.items():

        total_val = ws.cell(row=total_row, column=col).value

        if key == "Total Stones":
            ws.cell(row=ratio_total_row, column=col).value = total_val
        else:
            ratio = (total_val / total_stones) if total_stones != 0 else 0
            cell = ws.cell(row=ratio_total_row, column=col)
            cell.value = ratio
            cell.number_format = '0%'

    # ===== SAVE FILE =====
    output_name = f"Final_Inventory_Output_{name}.xlsx"
    wb.save( project_root/'Summary Tables'/f'{output_name}')

    print(f"✅ Saved: {output_name}")