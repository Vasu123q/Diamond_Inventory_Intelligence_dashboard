from openpyxl import load_workbook, Workbook
from copy import copy
from openpyxl.utils.dataframe import dataframe_to_rows
import pandas as pd
from pathlib import Path

project_root = Path(__file__).resolve().parent

wb_l1 = load_workbook(project_root/'Summary Tables'/'Final_Inventory_Output_L1.xlsx')
wb_l2 = load_workbook(project_root/'Summary Tables'/'Final_Inventory_Output_L2.xlsx')
wb_l3 = load_workbook(project_root/'Summary Tables'/'Final_Inventory_Output_L3.xlsx')

l1 = pd.read_excel(project_root/'L1'/'ML_Prediction_Level1.xlsx')
l2 = pd.read_excel(project_root/'L2'/'ML_Prediction_Level2.xlsx')
l3 = pd.read_excel(project_root/'L3'/'ML_Prediction_Level3.xlsx')

final_wb = Workbook()
final_wb.remove(final_wb.active)

from openpyxl.styles import Font

def make_header_bold(ws):
    for cell in ws[1]:   # first row = header
        cell.font = Font(bold=True)


def copy_sheet(source_wb, target_wb, new_name):
    source_ws = source_wb.active
    target_ws = target_wb.create_sheet(title=new_name)

    for row in source_ws.rows:
        for cell in row:
            new_cell = target_ws[cell.coordinate]
            new_cell.value = cell.value

            if cell.has_style:
                new_cell.font = copy(cell.font)
                new_cell.border = copy(cell.border)
                new_cell.fill = copy(cell.fill)
                new_cell.number_format = cell.number_format
                new_cell.alignment = copy(cell.alignment)

    # Column width
    for col in source_ws.column_dimensions:
        target_ws.column_dimensions[col].width = source_ws.column_dimensions[col].width



ws_l1 = final_wb.create_sheet("Level_1")
for r in dataframe_to_rows(l1, index=False, header=True):
    ws_l1.append(r)
make_header_bold(ws_l1) 
copy_sheet(wb_l1, final_wb, "Inventory_Summary_L1")


ws_l2 = final_wb.create_sheet("Level_2")
for r in dataframe_to_rows(l2, index=False, header=True):
    ws_l2.append(r)
make_header_bold(ws_l2) 
copy_sheet(wb_l2, final_wb, "Inventory_Summary_L2")


ws_l3 = final_wb.create_sheet("Level_3")
for r in dataframe_to_rows(l3, index=False, header=True):
    ws_l3.append(r)
make_header_bold(ws_l3) 
copy_sheet(wb_l3, final_wb, "Inventory_Summary_L3")


final_wb.save(project_root/"Final_Predicted.xlsx")

print("Concatation Completed")